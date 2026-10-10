# -*- coding: utf-8 -*-
"""Builds the tracker spreadsheet (XLSX + CSV) and the dashboard figures."""
import csv, json, re
from datetime import date
from collections import Counter, defaultdict

from dates import TODAY, days_late
from engine import (resolve, find_coordinator, find_members, member_list,
                    find_pay, ministries, government, sig_date, iso, nd)
from judgments import J
from pt import PT

ALLOWED = ["Concluído no prazo", "Concluído com atraso", "Em curso – dentro do prazo",
           "Em atraso", "Prorrogado", "Extinto / substituído", "Sem informação pública"]

acts = json.load(open("data/acts.json"))
acts.sort(key=lambda r: (r["dp"], r["num"]))

rows, problems = [], []

for i, r in enumerate(acts, start=1):
    j = J[r["num"]]
    nome_curto, missao_pt = PT[r["num"]]
    pub = iso(r["dp"])
    sig = sig_date(r["sig"], pub)
    gid = "GT-%d-%03d" % (pub.year, i)

    # ---- deadlines -------------------------------------------------
    dls = []
    for (label, kind, value, texto) in j["dl"]:
        due, how = resolve(kind, value, pub, sig)
        dls.append(dict(label=label, due=due, how=how, texto=texto))
    dated = [d for d in dls if d["due"] is not None]
    if j.get("no_primary") or not dated:
        primary = None
    else:
        pidx = j.get("primary", len(dls) - 1)
        primary = dls[pidx] if dls[pidx]["due"] else dated[-1]

    delivered = iso(j["delivered"]) if j.get("delivered") else None
    late = days_late(primary["due"], delivered) if primary else None
    if late is not None and late < 0:
        late = 0 if delivered else None  # delivered early, or not yet due

    status = j["status"]
    if status not in ALLOWED:
        problems.append(f"{r['num']}: status inválido '{status}'")

    # consistency checks
    if primary:
        overdue = primary["due"] < TODAY
        if status == "Em atraso" and not overdue:
            problems.append(f"{r['num']}: 'Em atraso' mas prazo {primary['due']} ainda não venceu")
        if status == "Em curso – dentro do prazo" and overdue:
            problems.append(f"{r['num']}: 'Em curso' mas prazo {primary['due']} já venceu")
        if status == "Concluído no prazo" and delivered and delivered > primary["due"]:
            problems.append(f"{r['num']}: 'no prazo' mas entrega {delivered} > prazo {primary['due']}")
        if status == "Concluído com atraso" and delivered and delivered <= primary["due"]:
            problems.append(f"{r['num']}: 'com atraso' mas entrega {delivered} <= prazo {primary['due']}")
    else:
        if status in ("Em atraso", "Prorrogado"):
            problems.append(f"{r['num']}: estado '{status}' sem prazo definido")
        if status == "Em curso – dentro do prazo" and not j.get("open_ended"):
            problems.append(f"{r['num']}: 'Em curso' sem prazo e sem open_ended")

    # ---- derived fields --------------------------------------------
    coord = find_coordinator(r["paras"]) or "Não identificado no ato"
    members = find_members(r["paras"])
    mlist = member_list(r["paras"])
    pay, pay_q = find_pay(r["paras"])

    mission_pt = ""
    for p in r["paras"]:
        if re.search(r"e criad[oa]|e constituid[oa]|tem (como|por) missao|a criacao de um grupo",
                     nd(p)):
            mission_pt = p
            break
    if not mission_pt:
        mission_pt = r["sum"]

    objectives = ""
    for p in r["paras"]:
        if re.search(r"objetivos|compete ao|tem como missao:|incumbe", nd(p)) and "|" in p:
            objectives = p
            break
    if not objectives:
        objectives = mission_pt

    xref = " ; ".join(r["xref"]) if r["xref"] else ""

    rows.append({
        "ID": gid,
        "Nome do grupo (sumário do DR)": r["sum"],
        "Designação curta": nome_curto,
        "Ato criador": f"{r['tipo']} n.º {r['num']}",
        "Link DR": r["url"],
        "Data de assinatura": sig.isoformat() if sig else "n/d",
        "Data de publicação (DR)": pub.isoformat(),
        "Entrada em vigor / produção de efeitos":
            next((p.split("| ", 1)[-1][:160] for p in r["paras"]
                  if re.search(r"produz efeitos|entra em vigor", nd(p))), "n/d"),
        "Governo": government(pub),
        "Ministérios / Secretarias de Estado": ministries(r["emit"]),
        "Emitente (completo)": r["emit"],
        "Coordenador": coord,
        "Membros (entidades)": mlist or (members[:600] or "Ver ato — composição não estruturada"),
        "Composição (texto do ato)": members[:1500],
        "Área de política": j["area"],
        "Missão (resumo)": missao_pt,
        "Missão (texto do ato)": mission_pt[:900],
        "Objetivos específicos (do ato)": objectives[:1800],
        "N.º de prazos no ato": len(dls),
        "Prazos — todos": " || ".join(
            f"{d['label']}: {d['due'].isoformat() if d['due'] else 'sem data'} "
            f"[{d['how']}] — «{d['texto'][:150]}»" for d in dls
        ) or "Sem prazo fixado no ato",
        "Prazo principal": primary["label"] if primary else "Sem prazo de entrega no ato",
        "Data-limite principal": primary["due"].isoformat() if primary else "Sem data no ato",
        "Como foi calculada": primary["how"] if primary else "—",
        "Participação remunerada": pay,
        "Remuneração (texto do ato)": pay_q[:300],
        "Estado": status,
        "Data de entrega / diploma": delivered.isoformat() if delivered else "",
        "Dias de atraso": late if late is not None else "",
        "Evidência": j["ev"],
        "Atos posteriores que citam este (DR)": xref,
        "Notas": j["notes"],
    })

# ------------------------------------------------------------------ outputs
cols = list(rows[0].keys())
with open("../grupos_de_trabalho.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=cols)
    w.writeheader()
    w.writerows(rows)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Grupos de trabalho"
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for rec in rows:
        ws.append([rec[c] for c in cols])

    STATUS_FILL = {
        "Concluído no prazo": "C6EFCE", "Concluído com atraso": "FFEB9C",
        "Em curso – dentro do prazo": "DDEBF7", "Em atraso": "FFC7CE",
        "Prorrogado": "FFE4C4", "Extinto / substituído": "E4DFEC",
        "Sem informação pública": "EDEDED",
    }
    sidx = cols.index("Estado") + 1
    for ri in range(2, len(rows) + 2):
        v = ws.cell(row=ri, column=sidx).value
        if v in STATUS_FILL:
            ws.cell(row=ri, column=sidx).fill = PatternFill("solid", fgColor=STATUS_FILL[v])
    widths = {"ID": 12, "Nome do grupo (sumário do DR)": 60, "Designação curta": 44, "Ato criador": 24,
              "Link DR": 54, "Estado": 26, "Evidência": 70, "Notas": 70,
              "Prazos — todos": 70, "Objetivos específicos (do ato)": 70,
              "Missão (resumo)": 60, "Missão (texto do ato)": 60, "Composição (texto do ato)": 60,
              "Membros (entidades)": 40, "Coordenador": 45, "Como foi calculada": 55,
              "Atos posteriores que citam este (DR)": 70, "Emitente (completo)": 55,
              "Ministérios / Secretarias de Estado": 40, "Remuneração (texto do ato)": 45,
              "Entrada em vigor / produção de efeitos": 40}
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 18)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows)+1}"
    wb.save("../grupos_de_trabalho.xlsx")
    xlsx_ok = True
except Exception as e:  # pragma: no cover
    xlsx_ok = False
    print("XLSX falhou:", e)

# ------------------------------------------------------------------ stats
stats = {}
stats["total"] = len(rows)
stats["by_gov"] = Counter(r["Governo"] for r in rows)
stats["by_status"] = Counter(r["Estado"] for r in rows)
stats["by_area"] = Counter(r["Área de política"] for r in rows)
stats["by_ministry"] = Counter(r["Ministérios / Secretarias de Estado"] for r in rows)
stats["by_pay"] = Counter(r["Participação remunerada"] for r in rows)
stats["by_month"] = Counter(r["Data de publicação (DR)"][:7] for r in rows)

delays = [r["Dias de atraso"] for r in rows if isinstance(r["Dias de atraso"], int) and r["Dias de atraso"] > 0]
stats["n_with_delay"] = len(delays)
stats["avg_delay"] = round(sum(delays) / len(delays), 1) if delays else 0
stats["median_delay"] = sorted(delays)[len(delays) // 2] if delays else 0
stats["max_delay"] = max(delays) if delays else 0
no_info = stats["by_status"]["Sem informação pública"]
stats["pct_no_info"] = round(100 * no_info / len(rows), 1)
undelivered = stats["by_status"]["Em atraso"]
stats["pct_overdue"] = round(100 * undelivered / len(rows), 1)
stats["pct_concluded"] = round(100 * (stats["by_status"]["Concluído no prazo"] +
                                      stats["by_status"]["Concluído com atraso"]) / len(rows), 1)
stats["no_deadline"] = sum(1 for r in rows if r["Data-limite principal"] == "Sem data no ato")

top = sorted((r for r in rows if isinstance(r["Dias de atraso"], int) and r["Dias de atraso"] > 0),
             key=lambda r: -r["Dias de atraso"])[:10]
stats["top10"] = [(r["ID"], r["Ato criador"], r["Nome do grupo (sumário do DR)"][:90],
                   r["Área de política"], r["Data-limite principal"], r["Dias de atraso"],
                   r["Estado"]) for r in top]

json.dump({k: (dict(v) if isinstance(v, Counter) else v) for k, v in stats.items()},
          open("data/stats.json", "w"), ensure_ascii=False, indent=1, default=str)
json.dump(rows, open("data/rows.json", "w"), ensure_ascii=False, indent=1, default=str)

print(f"Linhas: {len(rows)} | XLSX: {xlsx_ok}")
print("Por estado:", dict(stats["by_status"]))
print("Por governo:", dict(stats["by_gov"]))
print(f"Atrasos: n={stats['n_with_delay']} média={stats['avg_delay']} "
      f"mediana={stats['median_delay']} máx={stats['max_delay']}")
print(f"Sem prazo no ato: {stats['no_deadline']}")
if problems:
    print("\n!! INCONSISTÊNCIAS (%d):" % len(problems))
    for p in problems:
        print("  -", p)
else:
    print("\nSem inconsistências.")
